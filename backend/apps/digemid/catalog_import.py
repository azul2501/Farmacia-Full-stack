"""Importa el Catalogo de Productos Farmaceuticos de DIGEMID (Excel .xlsx o CSV) al catalogo maestro.

Las columnas se reconocen por su nombre, sin importar mayusculas, tildes ni separadores, para
tolerar cambios menores en el archivo que publica DIGEMID. Si una columna obligatoria no aparece,
el comando lo dice y no carga nada.
"""

import csv
import re
import unicodedata
from collections.abc import Iterator
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from apps.digemid.models import MasterProduct

# Campo del catalogo maestro -> nombres posibles de la columna (ya normalizados).
COLUMN_ALIASES = {
    "digemid_code": ["codprod", "codigoproducto", "codigo"],
    "name": ["nomprod", "nombreproducto", "producto", "nombre"],
    "concentration": ["concent", "concentracion"],
    "dosage_form": ["nomformfarm", "formafarmaceutica", "nomformfarmsimplif", "forma"],
    "presentation": ["presentac", "presentacion"],
    "fractions": ["fracciones", "fraccion"],
    "sanitary_registration": ["numregsan", "registrosanitario", "regsan", "nroregistrosanitario"],
    "holder": ["nomtitular", "titular"],
    "manufacturer": ["nomfabricante", "fabricante", "laboratorio"],
    "active_ingredient": ["nomifa", "principioactivo", "ifa", "dci"],
    "status": ["situacion", "estado", "condicion"],
    # Columnas propias de la lista base (no vienen en el Excel de DIGEMID).
    "suggested_price": ["preciosugerido", "precio"],
    "product_type": ["tipo", "tipoproducto"],
    "source": ["fuente", "origen"],
}
REQUIRED_FIELDS = ("digemid_code", "name")
TEXT_FIELDS = [name for name in COLUMN_ALIASES if name != "suggested_price"]
MAX_LENGTHS = {field_name: MasterProduct._meta.get_field(field_name).max_length for field_name in TEXT_FIELDS}
PRODUCT_TYPES = {"MEDICINE", "SUPPLY"}
BATCH_SIZE = 2000


def normalize_header(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]", "", text.lower())


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_price(value) -> Decimal | None:
    text = clean(value).replace("S/", "").replace(",", ".").strip()
    if not text:
        return None
    try:
        price = Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None
    return price if price > 0 else None


def _values(row: list, mapping: dict[str, int]) -> dict:
    values = {}
    for field_name, index in mapping.items():
        raw = row[index] if index < len(row) else ""
        if field_name == "suggested_price":
            values[field_name] = parse_price(raw)
        else:
            values[field_name] = clean(raw)[: MAX_LENGTHS[field_name]]
    if "product_type" in values:
        kind = values["product_type"].upper()
        values["product_type"] = kind if kind in PRODUCT_TYPES else "MEDICINE"
    return values


@dataclass
class ImportResult:
    mapping: dict[str, str] = field(default_factory=dict)
    missing: list[str] = field(default_factory=list)
    rows_read: int = 0
    rows_skipped: int = 0
    rows_saved: int = 0


def _rows(path: Path) -> Iterator[list]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        for row in sheet.iter_rows(values_only=True):
            yield list(row)
        workbook.close()
        return
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    lines = text.splitlines()
    # El delimitador mas frecuente en las primeras lineas (DIGEMID suele usar ";").
    sample = "\n".join(lines[:20])
    delimiter = max(",;\t|", key=sample.count)
    yield from csv.reader(lines, delimiter=delimiter)


def _find_header(rows: Iterator[list]) -> tuple[dict[str, int], list[str], list[str]]:
    """Busca la fila de encabezados entre las primeras 20 filas."""
    for _ in range(20):
        row = next(rows, None)
        if row is None:
            break
        normalized = [normalize_header(cell) for cell in row]
        mapping = {}
        for field_name, aliases in COLUMN_ALIASES.items():
            for alias in aliases:
                if alias in normalized:
                    mapping[field_name] = normalized.index(alias)
                    break
        if all(name in mapping for name in REQUIRED_FIELDS):
            headers = [clean(cell) for cell in row]
            missing = [name for name in COLUMN_ALIASES if name not in mapping]
            return mapping, headers, missing
    return {}, [], list(REQUIRED_FIELDS)


def import_catalog(path: Path, *, dry_run: bool = False, default_source: str = "DIGEMID") -> ImportResult:
    result = ImportResult()
    rows = _rows(path)
    mapping, headers, missing = _find_header(rows)
    if not mapping:
        result.missing = missing
        return result
    result.mapping = {field_name: headers[index] for field_name, index in mapping.items()}
    result.missing = missing

    batch: dict[str, MasterProduct] = {}
    update_fields = [name for name in COLUMN_ALIASES if name != "digemid_code" and name in mapping]
    if "source" not in mapping:
        update_fields.append("source")

    def flush():
        if batch and not dry_run:
            MasterProduct.objects.bulk_create(
                list(batch.values()),
                update_conflicts=True,
                unique_fields=["digemid_code"],
                update_fields=[*update_fields, "updated_at"],
            )
        result.rows_saved += len(batch)
        batch.clear()

    for row in rows:
        result.rows_read += 1
        values = _values(row, mapping)
        values.setdefault("source", default_source)
        if not values.get("digemid_code") or not values.get("name"):
            result.rows_skipped += 1
            continue
        batch[values["digemid_code"]] = MasterProduct(**values)
        if len(batch) >= BATCH_SIZE:
            flush()
    flush()
    return result
